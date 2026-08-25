#!/usr/bin/env python3
"""Extract SPARTA's Space Vehicle Logging Best Practices workbook to JSON (AINOS3-95).

The workbook (`data/sparta/Space_Vehicle_Logging_Best_Practices-Distro_A.xlsx`, from
<https://sparta.aerospace.org/related-work/iob>) is git-ignored — Aerospace Corporation's
to distribute — so every downstream artifact has to be re-derivable from a fresh download.
This is that derivation step: 19 sheets in, one JSON out, no interpretation.

Two shapes are recognised:

* **subsystem sheets** — a two-line preamble, then a header row whose columns are
  `Subsystem / Subcomponent / SPARTA ID Ref / ... / Minimum Logged Data / ...`. Column
  titles vary by whitespace and line-break between sheets, so headers are normalised.
* **`SPARTA_Mapping`** — the technique to log-source index, `ID / Name / Description /
  Subsystem / Sub-Component / Data Type / Attack Vector`.

Anything else is emitted with its raw rows and marked `reference`.

Run:
    python3 components/onair/training/extract_logging_workbook.py \\
        [--workbook data/sparta/Space_Vehicle_Logging_Best_Practices-Distro_A.xlsx] \\
        [--json data/sparta/logging_workbook.json]
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import warnings

DEFAULT_WORKBOOK = "data/sparta/Space_Vehicle_Logging_Best_Practices-Distro_A.xlsx"
DEFAULT_JSON = "data/sparta/logging_workbook.json"

# Sheets that carry per-row logging recommendations, in workbook order.
SUBSYSTEM_SHEETS = [
    "Propulsion", "ADCS", "EPS", "GN&C", "C&DH", "TT&C", "SMS", "TCS",
    "Payload - Imagery", "Payload - RF", "Payload - OCT",
    "Payload - Data Processing", "Payload - Hosted",
]
MAPPING_SHEET = "SPARTA_Mapping"

# `EX-0012.01,  EX-0012.10` / `IA-0003, IA-0008.02` -> ['EX-0012.01', ...]
_ID_RE = re.compile(r"\b([A-Z]{2,3}-\d{4}(?:\.\d{2})?)\b")


def norm_header(s) -> str:
    """`Impact/ Prioritization\n(Low/Medium/High)` -> `impact/prioritization`."""
    if s is None:
        return ""
    s = re.sub(r"\s+", " ", str(s)).strip().lower()
    s = re.sub(r"\s*\(.*?\)\s*", "", s)          # drop parenthetical legends
    return re.sub(r"\s*/\s*", "/", s).strip()


def clean(v):
    if v is None:
        return ""
    return re.sub(r"[ \t]*\n[ \t]*", " ", str(v)).strip()


def find_header_row(rows, want):
    """First row containing every token in `want` (normalised)."""
    for i, row in enumerate(rows):
        norm = {norm_header(c) for c in row}
        if all(w in norm for w in want):
            return i
    return None


def parse_table(rows, want):
    hdr_i = find_header_row(rows, want)
    if hdr_i is None:
        return None, None
    header = [norm_header(c) for c in rows[hdr_i]]
    out = []
    for row in rows[hdr_i + 1:]:
        rec = {h: clean(v) for h, v in zip(header, row) if h}
        if not any(rec.values()):
            continue
        out.append(rec)
    return header, out


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--workbook", default=DEFAULT_WORKBOOK)
    ap.add_argument("--json", default=DEFAULT_JSON)
    args = ap.parse_args(argv)

    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl is required: pip install openpyxl")

    with warnings.catch_warnings():           # print-area defined names, harmless
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(args.workbook, data_only=True)

    sha = hashlib.sha256(open(args.workbook, "rb").read()).hexdigest()
    doc = {"workbook": args.workbook, "sha256": sha, "sheets": {}}

    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        entry = {"rows": ws.max_row, "cols": ws.max_column}
        if ws.title in SUBSYSTEM_SHEETS:
            header, recs = parse_table(rows, ["subsystem", "sparta id ref",
                                              "minimum logged data"])
            for r in recs or []:
                r["sparta_ids"] = sorted(set(_ID_RE.findall(r.get("sparta id ref", ""))))
            entry.update(kind="subsystem", header=header, records=recs or [])
        elif ws.title == MAPPING_SHEET:
            header, recs = parse_table(rows, ["id", "name", "data type"])
            for r in recs or []:
                r["sparta_ids"] = sorted(set(_ID_RE.findall(r.get("id", ""))))
            entry.update(kind="mapping", header=header, records=recs or [])
        else:
            entry.update(kind="reference",
                         raw=[[clean(c) for c in row] for row in rows
                              if any(c is not None for c in row)])
        doc["sheets"][ws.title] = entry

    with open(args.json, "w") as f:
        json.dump(doc, f, indent=1)

    for name, e in doc["sheets"].items():
        n = len(e.get("records", e.get("raw", [])))
        print(f"{name:30s} {e['kind']:10s} {e['rows']:3d}x{e['cols']:<3d} -> {n} rows")
    total = sum(len(e.get("records", [])) for e in doc["sheets"].values())
    print(f"\n{total} recommendation rows -> {args.json}  (sha256 {sha[:12]})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
